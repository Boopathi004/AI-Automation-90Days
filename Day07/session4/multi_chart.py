from openpyxl import Workbook #type:ignore
from openpyxl.chart import BarChart, PieChart, Reference #type:ignore

wb = Workbook()

sheet = wb.active

sheet.append(["Department", "Salary"])

sheet.append(["IT",80000])
sheet.append(["HR",45000])
sheet.append(["Sales",55000])
sheet.append(["Admin",35000])

data = Reference(sheet, min_col=2, min_row=1, max_row=5)
labels = Reference(sheet, min_col=1, min_row=2, max_row=5)

bar = BarChart()
bar.title = "Salary"

bar.add_data(data, titles_from_data=True)
bar.set_categories(labels)

pie = PieChart()
pie.title = "Distribution"

pie.add_data(data, titles_from_data=True)
pie.set_categories(labels)

sheet.add_chart(bar, "D2")
sheet.add_chart(pie, "D18")

wb.save("multiple_charts.xlsx")